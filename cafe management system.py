import tkinter as tk
from tkinter import ttk,messagebox
from openpyxl import Workbook,load_workbook
from pathlib import Path
from datetime import datetime
import os
F=Path(__file__).parent/"cafe_records.xlsx"
MENU={"Tea":20,"Coffee":40,"Cold Coffee":80,"Sandwich":70,"Burger":100,"Pizza":150,"French Fries":80,"Pasta":120,"Cold Drink":50}
H=["ID","Customer","Item","Quantity","Price","Total","Date"]
def excel():
    if not F.exists():
        w=Workbook();w.active.append(H);w.save(F);w.close()
def rows():
    excel();w=load_workbook(F);r=list(w.active.iter_rows(min_row=2,values_only=True));w.close();return r
def save(r):
    w=load_workbook(F);w.active.append(r);w.save(F);w.close();os.startfile(F)
def app():
    root=tk.Tk();root.title("Cafe Management System");root.geometry("1000x600")
    clear=lambda:[x.destroy() for x in root.winfo_children()]
    def login():
        clear();f=tk.Frame(root,bg="white",padx=40,pady=30);f.place(relx=.5,rely=.5,anchor="center")
        tk.Label(f,text="CAFE MANAGEMENT SYSTEM",font=("Arial",20,"bold"),fg="darkblue",bg="white").pack(pady=10)
        tk.Label(f,text="Username",bg="white").pack();u=tk.Entry(f,width=30);u.pack(pady=5)
        tk.Label(f,text="Password",bg="white").pack();p=tk.Entry(f,width=30,show="*");p.pack(pady=5)
        check=lambda: messagebox.showerror("Error","Enter username and password") if not u.get().strip() or not p.get().strip() else dashboard()
        tk.Button(f,text="LOGIN",width=25,bg="darkblue",fg="white",command=check).pack(pady=10)
    def logout():
        if messagebox.askyesno("Logout","Do you want to logout?"):login()
    def dashboard():
        clear();tk.Label(root,text="CAFE MANAGEMENT SYSTEM",font=("Arial",22,"bold"),bg="lightblue",fg="darkblue").pack(fill="x",pady=20);[tk.Button(root,text=text,width=20,height=2,command=cmd).pack(pady=8) for text,cmd in [("Add Record",add),("View/Search",view),("Logout",logout)]];r=rows();tk.Label(root,text=f"Records: {len(r)} | Sales: Rs.{sum(x[5] for x in r):.2f}",bg="lightblue",font=("Arial",13)).pack(pady=20)
    def add():
        clear();tk.Label(root,text="ADD RECORD",font=("Arial",20,"bold"),bg="lightblue").pack(pady=15);f=tk.Frame(root,bg="white",padx=25,pady=15);f.pack()
        tk.Label(f,text="Customer",bg="white").grid(row=0,column=0,padx=8,pady=8);c=tk.Entry(f,width=28);c.grid(row=0,column=1)
        tk.Label(f,text="Item",bg="white").grid(row=1,column=0,padx=8,pady=8);i=ttk.Combobox(f,values=list(MENU),state="readonly",width=25);i.grid(row=1,column=1)
        tk.Label(f,text="Quantity",bg="white").grid(row=2,column=0,padx=8,pady=8);q=tk.Entry(f,width=28);q.grid(row=2,column=1)
        price=tk.Label(f,text="Price: Rs. 0",bg="white");price.grid(row=3,column=1);i.bind("<<ComboboxSelected>>",lambda e:price.config(text=f"Price: Rs. {MENU[i.get()]}"))
        def adddata():
            if not c.get().strip() or not i.get() or not q.get().isdigit() or int(q.get())<=0:return messagebox.showerror("Error","Enter valid customer, item and quantity")
            n=int(q.get());pr=MENU[i.get()];save([len(rows())+1,c.get(),i.get(),n,pr,pr*n,datetime.now().strftime("%d-%m-%Y %H:%M")]);messagebox.showinfo("Success","Saved and Excel opened");c.delete(0,"end");q.delete(0,"end")
        tk.Button(f,text="SAVE",bg="green",fg="white",command=adddata).grid(row=4,column=0,pady=12);tk.Button(f,text="CLEAR/RESET",command=lambda:(c.delete(0,"end"),i.set(""),q.delete(0,"end"),price.config(text="Price: Rs. 0"))).grid(row=4,column=1)
        tk.Button(root,text="Dashboard",command=dashboard).pack(pady=8)
    def view():
        clear();tk.Label(root,text="CAFE RECORDS",font=("Arial",20,"bold"),bg="lightblue").pack(pady=8);f=tk.Frame(root,bg="lightblue");f.pack()
        s=tk.Entry(f,width=25);s.pack(side="left");t=ttk.Treeview(root,columns=H,show="headings")
        for x in H:t.heading(x,text=x);t.column(x,width=110)
        t.pack(fill="both",expand=True,padx=8,pady=8)
        load=lambda d=None:(t.delete(*t.get_children()),[t.insert("",tk.END,values=r) for r in (rows() if d is None else d)])
        def search():
            x=s.get().lower();load([r for r in rows() if x in str(r[1]).lower() or x in str(r[2]).lower()])
        def update():
            z=t.selection()
            if not z:return messagebox.showwarning("Warning","Select a record")
            v=t.item(z[0],"values");w=tk.Toplevel(root);w.title("Update");w.geometry("350x300")
            c=tk.Entry(w);c.pack(pady=5);c.insert(0,v[1]);i=ttk.Combobox(w,values=list(MENU),state="readonly");i.pack(pady=5);i.set(v[2]);q=tk.Entry(w);q.pack(pady=5);q.insert(0,v[3])
            def up():
                if not c.get().strip() or not i.get() or not q.get().isdigit() or int(q.get())<=0:return messagebox.showerror("Error","Invalid data")
                n=int(q.get());pr=MENU[i.get()];wb=load_workbook(F);ws=wb.active
                for row in range(2,ws.max_row+1):
                    if ws.cell(row,1).value==int(v[0]):
                        for col,val in enumerate([int(v[0]),c.get(),i.get(),n,pr,pr*n,v[6]],1):ws.cell(row,col).value=val
                        break
                wb.save(F);wb.close();w.destroy();load()
            tk.Button(w,text="UPDATE",command=up).pack(pady=10)
        def delete():
            z=t.selection()
            if not z:return messagebox.showwarning("Warning","Select a record")
            v=t.item(z[0],"values")
            if messagebox.askyesno("Delete","Delete selected record?"):
                wb=load_workbook(F);ws=wb.active
                for row in range(2,ws.max_row+1):
                    if ws.cell(row,1).value==int(v[0]):ws.delete_rows(row);break
                wb.save(F);wb.close();load()
        for text,cmd in [("SEARCH",search),("SHOW ALL",lambda:load()),("UPDATE",update),("DELETE",delete),("LOGOUT",logout)]:tk.Button(f,text=text,command=cmd).pack(side="left",padx=3)
        load()
    excel();login();root.mainloop()
app()