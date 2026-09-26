import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from datetime import datetime
from pathlib import Path

FILE = "cafe_records.xlsx"
MENU = {"Tea": 20, "Coffee": 40, "Cold Coffee": 80, "Sandwich": 70,
        "Burger": 100, "Pizza": 150, "French Fries": 80, "Pasta": 120,
        "Cold Drink": 50}
HEADERS = ["ID", "Customer", "Item", "Quantity", "Price", "Total", "Date"]


def create_excel():
    if not Path(FILE).exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Cafe Records"
        ws.append(HEADERS)
        wb.save(FILE)
        wb.close()


def get_records():
    create_excel()
    wb = load_workbook(FILE)
    ws = wb.active
    data = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return data


def next_id():
    ids = [int(r[0]) for r in get_records() if r[0] is not None]
    return max(ids, default=0) + 1


def add_excel(data):
    wb = load_workbook(FILE)
    wb.active.append(data)
    wb.save(FILE)
    wb.close()


def update_excel(record_id, data):
    wb = load_workbook(FILE)
    ws = wb.active
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == record_id:
            for col, value in enumerate(data, 1):
                ws.cell(row, col).value = value
            wb.save(FILE)
            wb.close()
            return True
    wb.close()
    return False


def delete_excel(record_id):
    wb = load_workbook(FILE)
    ws = wb.active
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == record_id:
            ws.delete_rows(row, 1)
            wb.save(FILE)
            wb.close()
            return True
    wb.close()
    return False


class CafeApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Cafe Management System")
        self.root.geometry("1100x650")
        self.root.configure(bg="lightblue")
        create_excel()
        self.login_page()

    def clear(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    # ---------- LOGIN + VALIDATION ----------
    def login_page(self):
        self.clear()
        box = tk.Frame(self.root, bg="white", padx=40, pady=30)
        box.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(box, text="CAFE MANAGEMENT SYSTEM", font=("Arial", 22, "bold"),
                 fg="darkblue", bg="white").pack(pady=15)
        tk.Label(box, text="Username", bg="white").pack(anchor="w")
        self.username = tk.Entry(box, width=35)
        self.username.pack(pady=5)
        tk.Label(box, text="Password", bg="white").pack(anchor="w")
        self.password = tk.Entry(box, width=35, show="*")
        self.password.pack(pady=5)
        tk.Button(box, text="LOGIN", width=30, bg="darkblue", fg="white",
                  command=self.login).pack(pady=15)
        tk.Label(box, text="Enter any username and password", fg="gray",
                 bg="white").pack()
        self.username.focus()
        self.root.bind("<Return>", lambda e: self.login())

    def login(self):
        # Validation only checks that both fields are filled.
        # There is no fixed username/password.
        if not self.username.get().strip():
            messagebox.showerror("Login Error", "Please enter username.")
            return
        if not self.password.get():
            messagebox.showerror("Login Error", "Please enter password.")
            return
        self.root.unbind("<Return>")
        self.dashboard()

    # ---------- DASHBOARD ----------
    def dashboard(self):
        self.clear()
        header = tk.Frame(self.root, bg="darkblue", height=65)
        header.pack(fill="x")
        tk.Label(header, text="CAFE MANAGEMENT SYSTEM", font=("Arial", 21, "bold"),
                 fg="white", bg="darkblue").pack(side="left", padx=20, pady=15)
        tk.Button(header, text="Logout", bg="red", fg="white",
                  command=self.logout).pack(side="right", padx=20)

        tk.Label(self.root, text="Dashboard", font=("Arial", 20, "bold"),
                 bg="lightblue", fg="darkblue").pack(pady=20)

        menu_frame = tk.Frame(self.root, bg="white", padx=20, pady=12)
        menu_frame.pack(pady=5)
        tk.Label(menu_frame, text="Cafe Menu", font=("Arial", 14, "bold"),
                 bg="white", fg="darkblue").grid(row=0, column=0, columnspan=3, pady=5)
        for i, (item, price) in enumerate(MENU.items()):
            tk.Label(menu_frame, text=f"{item} - Rs. {price}", bg="white", width=18,
                     anchor="w").grid(row=i // 3 + 1, column=i % 3, padx=8, pady=2)

        buttons = tk.Frame(self.root, bg="lightblue")
        buttons.pack(pady=20)
        for text, command in [("Add Record", self.add_page),
                              ("View Records", self.view_page)]:
            tk.Button(buttons, text=text, width=20, height=2, bg="white",
                      font=("Arial", 11, "bold"), command=command).pack(side="left", padx=10)

        records = get_records()
        sales = sum(float(r[5]) for r in records if r[5] is not None)
        tk.Label(self.root, text=f"Total Records: {len(records)}    |    Total Sales: Rs. {sales:.2f}",
                 font=("Arial", 13, "bold"), bg="lightblue", fg="darkgreen").pack(pady=10)

    def logout(self):
        if messagebox.askyesno("Logout", "Do you want to logout?"):
            self.login_page()

    # ---------- ADD RECORD + SAVE + CALCULATION ----------
    def add_page(self):
        self.clear()
        tk.Label(self.root, text="ADD CAFE RECORD", font=("Arial", 20, "bold"),
                 bg="lightblue", fg="darkblue").pack(pady=15)
        form = tk.Frame(self.root, bg="white", padx=25, pady=20)
        form.pack()

        tk.Label(form, text="Customer Name", bg="white").grid(row=0, column=0, padx=10, pady=10)
        self.customer = tk.Entry(form, width=30)
        self.customer.grid(row=0, column=1, padx=10, pady=10)

        tk.Label(form, text="Food Item", bg="white").grid(row=1, column=0, padx=10, pady=10)
        self.item = ttk.Combobox(form, values=list(MENU), state="readonly", width=27)
        self.item.grid(row=1, column=1, padx=10, pady=10)
        self.item.bind("<<ComboboxSelected>>", self.show_price)

        tk.Label(form, text="Quantity", bg="white").grid(row=2, column=0, padx=10, pady=10)
        self.quantity = tk.Entry(form, width=30)
        self.quantity.grid(row=2, column=1, padx=10, pady=10)

        tk.Label(form, text="Price", bg="white").grid(row=3, column=0, padx=10, pady=10)
        self.price = tk.Entry(form, width=30, state="readonly")
        self.price.grid(row=3, column=1, padx=10, pady=10)

        tk.Label(form, text="Total", bg="white").grid(row=4, column=0, padx=10, pady=10)
        self.total = tk.Entry(form, width=30, state="readonly")
        self.total.grid(row=4, column=1, padx=10, pady=10)

        self.quantity.bind("<KeyRelease>", self.calculate_total)
        tk.Button(form, text="SAVE", width=15, bg="green", fg="white",
                  command=self.save_record).grid(row=5, column=0, pady=20)
        tk.Button(form, text="CLEAR / RESET", width=15,
                  command=self.reset_form).grid(row=5, column=1, pady=20)
        tk.Button(self.root, text="Back to Dashboard", command=self.dashboard).pack(pady=10)

    def show_price(self, event=None):
        self.set_readonly(self.price, MENU.get(self.item.get(), ""))
        self.calculate_total()

    def calculate_total(self, event=None):
        try:
            total = MENU[self.item.get()] * int(self.quantity.get())
            self.set_readonly(self.total, total)
        except (ValueError, KeyError):
            self.set_readonly(self.total, "")

    def set_readonly(self, entry, value):
        entry.config(state="normal")
        entry.delete(0, tk.END)
        entry.insert(0, value)
        entry.config(state="readonly")

    def save_record(self):
        customer = self.customer.get().strip()
        item = self.item.get()
        qty_text = self.quantity.get().strip()
        if not customer:
            messagebox.showerror("Error", "Please enter customer name.")
            return
        if not item:
            messagebox.showerror("Error", "Please select food item.")
            return
        try:
            qty = int(qty_text)
            if qty <= 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Quantity must be a positive number.")
            return

        price = MENU[item]
        add_excel([next_id(), customer, item, qty, price, price * qty,
                   datetime.now().strftime("%d-%m-%Y %H:%M:%S")])
        messagebox.showinfo("Success", "Record saved successfully.")
        self.reset_form()

    def reset_form(self):
        self.customer.delete(0, tk.END)
        self.item.set("")
        self.quantity.delete(0, tk.END)
        self.set_readonly(self.price, "")
        self.set_readonly(self.total, "")
        self.customer.focus()

    # ---------- VIEW + TREEVIEW + SEARCH ----------
    def view_page(self):
        self.clear()
        top = tk.Frame(self.root, bg="lightblue")
        top.pack(fill="x", pady=10)
        tk.Label(top, text="CAFE RECORDS", font=("Arial", 20, "bold"),
                 bg="lightblue", fg="darkblue").pack(side="left", padx=20)
        tk.Button(top, text="Dashboard", command=self.dashboard).pack(side="right", padx=8)
        tk.Button(top, text="Logout", bg="red", fg="white",
                  command=self.logout).pack(side="right", padx=8)

        search_frame = tk.Frame(self.root, bg="lightblue")
        search_frame.pack(pady=8)
        tk.Label(search_frame, text="Search Customer / Item:", bg="lightblue",
                 font=("Arial", 11, "bold")).pack(side="left")
        self.search_entry = tk.Entry(search_frame, width=30)
        self.search_entry.pack(side="left", padx=10)
        tk.Button(search_frame, text="SEARCH", command=self.search).pack(side="left", padx=4)
        tk.Button(search_frame, text="SHOW ALL", command=self.load_records).pack(side="left", padx=4)

        frame = tk.Frame(self.root)
        frame.pack(fill="both", expand=True, padx=15, pady=8)
        self.tree = ttk.Treeview(frame, columns=HEADERS, show="headings")
        for col in HEADERS:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=125, anchor="center")
        scroll_y = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.tree.pack(side="left", fill="both", expand=True)
        scroll_y.pack(side="right", fill="y")
        scroll_x.pack(side="bottom", fill="x")

        actions = tk.Frame(self.root, bg="lightblue")
        actions.pack(pady=10)
        tk.Button(actions, text="UPDATE", width=15, bg="orange",
                  command=self.update_selected).pack(side="left", padx=5)
        tk.Button(actions, text="DELETE", width=15, bg="red", fg="white",
                  command=self.delete_selected).pack(side="left", padx=5)
        tk.Button(actions, text="REFRESH", width=15,
                  command=self.load_records).pack(side="left", padx=5)
        self.tree.bind("<Double-1>", self.update_selected)
        self.load_records()

    def load_records(self, records=None):
        for row in self.tree.get_children():
            self.tree.delete(row)
        for row in get_records() if records is None else records:
            self.tree.insert("", tk.END, values=row)

    def search(self):
        text = self.search_entry.get().strip().lower()
        if not text:
            self.load_records()
            return
        result = [r for r in get_records()
                  if text in str(r[1]).lower() or text in str(r[2]).lower()]
        self.load_records(result)

    # ---------- UPDATE ----------
    def update_selected(self, event=None):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select a record.")
            return
        values = self.tree.item(selected[0], "values")
        record_id = int(values[0])

        win = tk.Toplevel(self.root)
        win.title("Update Record")
        win.geometry("430x360")
        win.configure(bg="lightblue")
        win.grab_set()

        tk.Label(win, text="UPDATE RECORD", font=("Arial", 18, "bold"),
                 bg="lightblue", fg="darkblue").pack(pady=15)
        form = tk.Frame(win, bg="white", padx=20, pady=15)
        form.pack()

        tk.Label(form, text="Customer Name", bg="white").grid(row=0, column=0, padx=8, pady=10)
        customer = tk.Entry(form, width=25)
        customer.grid(row=0, column=1, padx=8, pady=10)
        customer.insert(0, values[1])

        tk.Label(form, text="Food Item", bg="white").grid(row=1, column=0, padx=8, pady=10)
        item = ttk.Combobox(form, values=list(MENU), state="readonly", width=22)
        item.grid(row=1, column=1, padx=8, pady=10)
        item.set(values[2])

        tk.Label(form, text="Quantity", bg="white").grid(row=2, column=0, padx=8, pady=10)
        quantity = tk.Entry(form, width=25)
        quantity.grid(row=2, column=1, padx=8, pady=10)
        quantity.insert(0, values[3])

        def save_update():
            name, selected_item = customer.get().strip(), item.get()
            if not name or not selected_item:
                messagebox.showerror("Error", "Please fill all fields.", parent=win)
                return
            try:
                qty = int(quantity.get())
                if qty <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Error", "Quantity must be a positive number.", parent=win)
                return

            price = MENU[selected_item]
            data = [record_id, name, selected_item, qty, price, price * qty, values[6]]
            if update_excel(record_id, data):
                messagebox.showinfo("Success", "Record updated successfully.", parent=win)
                win.destroy()
                self.load_records()

        tk.Button(form, text="UPDATE", width=15, bg="green", fg="white",
                  command=save_update).grid(row=3, column=0, columnspan=2, pady=15)

    # ---------- DELETE ----------
    def delete_selected(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Warning", "Please select a record.")
            return
        values = self.tree.item(selected[0], "values")
        record_id = int(values[0])
        if messagebox.askyesno("Confirm Delete", f"Delete record ID {record_id}?"):
            if delete_excel(record_id):
                messagebox.showinfo("Success", "Record deleted successfully.")
                self.load_records()


root = tk.Tk()
CafeApp(root)
root.mainloop()
